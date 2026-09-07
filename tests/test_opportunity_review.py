import copy
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_workbook import _apply_plan_overrides, _lookup_live_price, _price_profile, build_workbook
from input_integrity import source_hashes
from opportunity_review import GATES, SCHEMA, baseline_for, evaluate_review, review_digest, template_review
from quote_plan import QuotePlanError, apply_quote_plan, build_quote_plan, validate_quote_plan
from parse_inputs import normalize_inputs
from run_pipeline import _render_executive_summary, main as pipeline_main


CONTEXT = {"region": "eastus", "currency": "USD", "scenario": "conservative"}
PRICE = {
  "source": "API",
  "best": {"retailPrice": 0.1, "unitOfMeasure": "1 Hour", "effectiveStartDate": "2026-01-01"},
  "links": ["https://prices.azure.com/api/retail/prices"], "note": "Test price",
}


class OpportunityReviewTests(unittest.TestCase):
  def setUp(self):
    self.directory = tempfile.TemporaryDirectory(dir=ROOT / "tests", prefix=".review-test-")
    self.addCleanup(self.directory.cleanup)
    self.root = Path(self.directory.name)
    self.specs_path = self.root / "specs.json"
    self.boq_path = self.root / "boq.csv"
    self.specs_path.write_text("{}", encoding="utf-8")
    self.boq_path.write_text("Service,Quantity\nEC2,730\n", encoding="utf-8")
    row = {
      "Service": "EC2", "Instance/SKU": "m5.xlarge", "Quantity": 730,
      "Source Unit": "Hours", "Capacity": "4 vCPU / 16 GB", "Monthly": 200,
      "Azure Service": "Virtual Machines", "Azure SKU": "Standard_D4s_v5",
    }
    self.normalized = {"aws_boq": [row, dict(row, **{"Instance/SKU": "m5.large"})],
                       "specs": {}, "missing_or_ambiguous": [],
                       "specs_source": str(self.specs_path), "aws_boq_source": str(self.boq_path),
                       "sourceBindings": source_hashes(self.specs_path, self.boq_path)}
    self.plan = build_quote_plan(self.normalized)
    self.review = template_review(self.plan, CONTEXT, "submission-ready", self.normalized)
    self.plan["opportunityReview"] = self.review
    self.review["baselineComplete"] = True
    self.review["documents"] = [{"id": "D1", "version": "1", "sourceReference": "rfp.pdf",
                                 "contentId": "customer-issued-revision-1"}]
    self.review["requirements"] = [{
      "id": "R1", "priority": "mandatory", "acceptanceCriterion": "Run in East US",
      "source": {"documentId": "D1", "location": "p. 4, section 2.1, table 1 row 3"},
      "qualification": "pass", "evidence": ["Architecture decision ADR-1 approved for East US"],
      "lineIds": [line["id"] for line in self.plan["lines"]],
    }]

  def approve(self, record):
    record["status"] = "approved"
    record["approval"] = {"name": "Seller Example", "confirmedAt": "2026-01-02T12:00:00Z",
                          "evidence": "User confirmation message review-123",
                          "reviewDigest": review_digest(self.review)}

  def approve_all(self):
    for record in self.review["assumptions"] + self.review["gates"]:
      self.approve(record)

  def assumption(self, material=True):
    result = {"id": "A1", "version": "1", "text": "730 runtime hours confirmed by customer",
              "material": material, "lineIds": [self.plan["lines"][0]["id"]],
              "status": "proposed", "owner": "Seller Example", "nextAction": "Confirm runtime hours"}
    self.review["assumptions"].append(result)
    return result

  def status(self, **kwargs):
    return evaluate_review(self.plan, CONTEXT, self.normalized, compiled=True, **kwargs)

  def compile(self, plan=None, name="quote.xlsx"):
    normalized = apply_quote_plan(self.normalized, plan or self.plan)
    with patch("build_workbook.resolve_best_price", return_value=PRICE):
      return build_workbook(normalized, CONTEXT, self.root / name, "conservative")

  def test_valid_review_and_independent_weighted_score(self):
    scored = copy.deepcopy(self.review["requirements"][0])
    scored.update(id="R2", priority="scored", score=0.8, weight=2)
    self.review["requirements"].append(scored)
    self.approve_all()
    status = self.status()
    self.assertTrue(status["releaseReady"])
    self.assertEqual("Release-ready", status["readiness"])
    self.assertEqual("Pass", status["eligibility"])
    self.assertEqual(0.8, status["weightedScore"])
    self.assertTrue(all(value == "approved" for value in status["approvals"].values()))

  def test_hard_fail_cannot_be_overridden_by_gates_or_score(self):
    self.review["requirements"][0]["qualification"] = "fail"
    scored = copy.deepcopy(self.review["requirements"][0])
    scored.update(id="R2", priority="scored", qualification="pass", score=1, weight=999)
    self.review["requirements"].append(scored)
    self.approve_all()
    status = self.status()
    self.assertEqual("Fail", status["eligibility"])
    self.assertEqual(1, status["weightedScore"])
    self.assertFalse(status["releaseReady"])
    self.assertFalse(status["qualifiedComparison"])

  def test_unresolved_and_incomplete_baselines_never_pass(self):
    for change in ("unresolved", "empty", "incomplete"):
      with self.subTest(change=change):
        plan = copy.deepcopy(self.plan)
        review = plan["opportunityReview"]
        if change == "unresolved":
          review["requirements"][0]["qualification"] = "unresolved"
          review["requirements"][0]["evidence"] = []
        elif change == "empty":
          review["requirements"] = []
        else:
          review["baselineComplete"] = False
        self.assertEqual("Unresolved", evaluate_review(plan, CONTEXT, self.normalized)["eligibility"])

  def test_pass_fail_require_evidence_and_precise_location(self):
    for field, value in (("evidence", []), ("source", {"documentId": "D1", "location": "  "})):
      with self.subTest(field=field):
        plan = copy.deepcopy(self.plan)
        plan["opportunityReview"]["requirements"][0][field] = value
        with self.assertRaises(QuotePlanError):
          validate_quote_plan(plan)

  def test_dangling_and_duplicate_references_rejected(self):
    mutations = [
      lambda r: r["requirements"][0]["source"].update(documentId="missing"),
      lambda r: r["requirements"][0].update(lineIds=["missing"]),
      lambda r: r["requirements"].append(copy.deepcopy(r["requirements"][0])),
      lambda r: r["documents"].append(copy.deepcopy(r["documents"][0])),
      lambda r: r["gates"].append(copy.deepcopy(r["gates"][0])),
      lambda r: r["requirements"][0].update(supersedes="missing"),
      lambda r: r["documents"][0].update(amends="missing"),
      lambda r: r["requirements"][0].update(lineIds=[self.plan["lines"][0]["id"]] * 2),
    ]
    for mutate in mutations:
      with self.subTest(mutate=mutate):
        plan = copy.deepcopy(self.plan)
        mutate(plan["opportunityReview"])
        with self.assertRaisesRegex(QuotePlanError, "duplicate|dangling"):
          validate_quote_plan(plan)

  def test_amended_requirement_explicitly_supersedes_old_fail(self):
    self.review["requirements"][0]["qualification"] = "fail"
    self.review["documents"].append({"id": "D2", "version": "2", "sourceReference": "amendment.pdf",
                                     "contentId": "customer-amendment-2", "amends": "D1"})
    replacement = copy.deepcopy(self.review["requirements"][0])
    replacement.update(id="R2", supersedes="R1", qualification="pass",
                       source={"documentId": "D2", "location": "p. 1, amended clause 2.1"})
    self.review["requirements"].append(replacement)
    self.approve_all()
    self.assertEqual("Pass", self.status()["eligibility"])
    self.assertEqual(1, self.status()["coverage"]["requirements"])

  def test_document_supersession_requires_explicit_requirement_replacement(self):
    self.review["documents"].append({"id": "D2", "version": "2", "sourceReference": "replacement.pdf",
                                     "contentId": "revision-2", "supersedes": "D1"})
    with self.assertRaisesRegex(QuotePlanError, "active requirement on superseded"):
      validate_quote_plan(self.plan)
    replacement = copy.deepcopy(self.review["requirements"][0])
    replacement.update(id="R2", supersedes="R1", source={"documentId": "D2", "location": "p. 1 clause 1"})
    self.review["requirements"].append(replacement)
    validate_quote_plan(self.plan)

  def test_cycles_and_ambiguous_replacements_rejected(self):
    for case in ("document-cycle", "requirement-cycle", "ambiguous"):
      with self.subTest(case=case):
        plan = copy.deepcopy(self.plan)
        review = plan["opportunityReview"]
        if case == "document-cycle":
          review["documents"][0]["amends"] = "D1"
        elif case == "requirement-cycle":
          review["requirements"][0]["supersedes"] = "R1"
        else:
          for key in ("R2", "R3"):
            review["requirements"].append(dict(review["requirements"][0], id=key, supersedes="R1"))
        with self.assertRaisesRegex(QuotePlanError, "cycle|ambiguous"):
          validate_quote_plan(plan)

  def test_material_assumption_blocks_only_affected_costs_and_reference_bypass(self):
    self.assumption()
    self.normalized["aws_boq"][0].update({"Meter Contains": "D4s v5", "Local Azure Total Cost": 1})
    self.review["baseline"] = baseline_for(self.plan, CONTEXT, self.normalized)
    normalized = apply_quote_plan(self.normalized, self.plan)
    self.assertTrue(normalized["aws_boq"][0]["Plan Validations"])
    self.assertFalse(normalized["aws_boq"][1]["Plan Validations"])
    summary = self.compile()
    self.assertEqual(1, summary["validateCount"])
    self.assertEqual("N/A ([VALIDATE])", summary["scenarioTotalsMonthly"]["conservative"])
    self.assertEqual(73, summary["scenarioPricedSubtotalsMonthly"]["conservative"]["amount"])
    self.assertEqual("Seller Example", summary["opportunityStatus"]["owner"])
    workbook = load_workbook(self.root / "quote.xlsx", data_only=True)
    self.addCleanup(workbook.close)
    rows = list(workbook["Mapping & Azure BOQ"].values)
    self.assertEqual("[VALIDATE]", rows[1][8])
    self.assertEqual(73, rows[2][8])

  def test_approved_material_assumption_allows_numeric_costs(self):
    self.assumption()
    self.approve_all()
    summary = self.compile()
    self.assertEqual(0, summary["validateCount"])
    self.assertEqual(146, summary["scenarioTotalsMonthly"]["conservative"])
    self.assertTrue(summary["opportunityStatus"]["releaseReady"])

  def test_nonmaterial_assumption_visible_without_blocking(self):
    self.assumption(material=False)
    summary = self.compile()
    self.assertEqual(0, summary["validateCount"])
    workbook = load_workbook(self.root / "quote.xlsx", data_only=True)
    self.addCleanup(workbook.close)
    self.assertIn("A1 v1", str(list(workbook["Assumptions & Levers"].values)))

  def test_catalog_pricing_cannot_bypass_material_blocker(self):
    profile = _price_profile("GitHub", "GitHub", "GitHub Team", "25 users", "GitHub Team")
    profile = _apply_plan_overrides(profile, {"Plan Validations": ["Unapproved material assumption"]})
    with patch("build_workbook.resolve_catalog_price") as catalog:
      result = _lookup_live_price("eastus", "USD", profile, "conservative", 25)
    catalog.assert_not_called()
    self.assertTrue(result["validate"])
    self.assertEqual(0, result["monthly"])

  def test_stale_plan_context_and_pricing_inputs_block_all_lines(self):
    self.assumption()
    self.approve_all()
    for case in ("quantity", "mapping", "region", "currency", "scenario", "meter"):
      with self.subTest(case=case):
        plan, context, normalized = copy.deepcopy(self.plan), dict(CONTEXT), copy.deepcopy(self.normalized)
        if case == "quantity":
          plan["lines"][0]["source"]["quantity"] = 999
        elif case == "mapping":
          plan["lines"][0]["target"]["sku"] = "Standard_D8s_v5"
        elif case == "meter":
          normalized["aws_boq"][0]["Meter Contains"] = "new meter"
        else:
          context[case] = {"region": "westus", "currency": "EUR", "scenario": "aggressive"}[case]
        status = evaluate_review(plan, context, normalized, compiled=True)
        self.assertFalse(status["releaseReady"])
        self.assertFalse(status["baselineCurrent"])
        self.assertEqual(2, len(status["blockedLineIds"]))
        with patch("build_workbook.resolve_best_price") as resolve:
          summary = build_workbook(apply_quote_plan(normalized, plan), context, self.root / "stale.xlsx", context["scenario"])
        resolve.assert_not_called()
        self.assertEqual(2, summary["validateCount"])

  def test_stale_assumption_version_and_requirement_approval(self):
    assumption = self.assumption()
    self.approve_all()
    assumption["version"] = "2"
    status = self.status()
    self.assertTrue(status["baselineCurrent"])
    self.assertEqual([self.plan["lines"][0]["id"]], status["blockedLineIds"])
    self.assertTrue(all(v == "stale" for v in status["approvals"].values()))
    self.approve_all()
    self.assertTrue(self.status()["releaseReady"])
    self.review["requirements"][0]["acceptanceCriterion"] = "Changed criterion"
    self.assertFalse(self.status()["releaseReady"])
    self.assertTrue(self.status()["blockedLineIds"])

  def test_changed_original_or_resumed_source_quantity_is_stale(self):
    self.assumption()
    self.approve_all()
    for resumed in (False, True):
      with self.subTest(resumed=resumed):
        normalized = apply_quote_plan(self.normalized, self.plan) if resumed else copy.deepcopy(self.normalized)
        normalized["aws_boq"][0]["Quantity"] = 999
        applied = apply_quote_plan(normalized, self.plan, CONTEXT)
        self.assertTrue(all(row["Plan Validations"] for row in applied["aws_boq"]))
        with patch("build_workbook.resolve_best_price") as resolve:
          summary = build_workbook(applied, CONTEXT, self.root / "source-changed.xlsx", "conservative")
        resolve.assert_not_called()
        self.assertFalse(summary["opportunityStatus"]["baselineCurrent"])
        self.assertEqual(2, summary["validateCount"])

  def test_optional_source_fields_use_identical_overlay_and_snapshot_semantics(self):
    for field in ("quantity", "unit", "capacity", "monthlyCost"):
      self.plan["lines"][0]["source"].pop(field)
    self.plan["lines"][0]["target"]["sku"] = "Standard_D8s_v5"
    self.review["baseline"] = baseline_for(self.plan, CONTEXT, self.normalized)
    self.approve_all()
    self.assertTrue(self.status()["baselineCurrent"])
    applied = apply_quote_plan(self.normalized, self.plan)
    self.assertEqual(730, applied["aws_boq"][0]["Quantity"])
    self.assertEqual(200, applied["aws_boq"][0]["Monthly"])
    self.assertEqual("Standard_D8s_v5", applied["aws_boq"][0]["Azure SKU"])
    self.assertEqual(self.review["baseline"], baseline_for(self.plan, CONTEXT, applied))
    reapplied = apply_quote_plan(applied, self.plan)
    self.assertEqual(applied["reviewSourceRows"], reapplied["reviewSourceRows"])
    self.assertFalse(reapplied["aws_boq"][0]["Plan Validations"])
    self.assertTrue(self.compile()["opportunityStatus"]["releaseReady"])
    for field, changed in (("Quantity", 999), ("Source Unit", "Months"), ("Capacity", "different"),
                           ("Monthly", 999), ("Azure SKU", "Standard_D16s_v5")):
      with self.subTest(field=field):
        edited = copy.deepcopy(applied)
        edited["aws_boq"][0][field] = changed
        status = evaluate_review(self.plan, CONTEXT, edited, compiled=True)
        self.assertFalse(status["baselineCurrent"])
        self.assertEqual(2, len(status["blockedLineIds"]))

  def test_normalization_records_source_content_hashes_and_unknown_costs(self):
    self.boq_path.write_text("Service,Quantity,Monthly\nEC2,1,\nEC2,1,0\nEC2,1,N/A\n", encoding="utf-8")
    normalized = normalize_inputs(self.specs_path, self.boq_path, self.root / "generated.json")
    self.assertEqual(source_hashes(self.specs_path, self.boq_path), normalized["sourceBindings"])
    self.assertEqual("", normalized["aws_boq"][0]["Monthly"])
    self.assertEqual(0, normalized["aws_boq"][1]["Monthly"])
    plan = build_quote_plan(normalized)
    self.assertNotIn("monthlyCost", plan["lines"][0]["source"])
    self.assertEqual(0, plan["lines"][1]["source"]["monthlyCost"])
    self.assertNotIn("monthlyCost", plan["lines"][2]["source"])

  def test_unknown_incumbent_costs_do_not_block_target_only_release(self):
    for provider in ("aws", "external"):
      for scenario in ("conservative", "compare-all"):
        with self.subTest(provider=provider, scenario=scenario):
          plan, normalized = copy.deepcopy(self.plan), copy.deepcopy(self.normalized)
          plan["sourceProvider"] = provider
          for row, line in zip(normalized["aws_boq"], plan["lines"]):
            row.pop("Monthly")
            if provider == "external":
              row["Monthly"] = "N/A"
            line["source"].pop("monthlyCost")
            line["source"]["provider"] = provider
          review = plan["opportunityReview"]
          context = dict(CONTEXT, scenario=scenario)
          review["baseline"] = baseline_for(plan, context, normalized)
          for gate in review["gates"]:
            gate["status"] = "approved"
            gate["approval"] = {"name": "Reviewer", "confirmedAt": "2026-01-02T12:00:00Z",
                                "evidence": "User message", "reviewDigest": review_digest(review)}
          applied = apply_quote_plan(normalized, plan, context)
          self.assertEqual("N/A" if provider == "external" else None, applied["aws_boq"][0]["Monthly"])
          with patch("build_workbook.resolve_best_price", return_value=PRICE):
            summary = build_workbook(applied, context, self.root / "unknown.xlsx", scenario)
          status = summary["opportunityStatus"]
          self.assertTrue(status["releaseReady"])
          self.assertFalse(status["qualifiedComparison"])
          self.assertFalse(status["incumbentBaselineComplete"])
          self.assertEqual(2, len(status["missingIncumbentCostLineIds"]))
          self.assertEqual("N/A (incumbent costs unknown)", summary["scenarioTotalsMonthly"]["aws"])
          self.assertEqual(146, summary["scenarioTotalsMonthly"]["conservative"])
          self.assertNotEqual("High", summary["decisionConfidence"])
          self.assertEqual("N/A", summary["recommendedScenario"])
          self.assertEqual([], summary["topDriversPositive"])
          self.assertEqual([], summary["topDriversNegative"])
          self.assertIsInstance(summary["scenarioSavings"]["conservative"], str)
          executive = _render_executive_summary(summary)
          self.assertIn("target-only estimate", executive)
          self.assertIn("incumbent costs unknown", executive)
          workbook = load_workbook(self.root / "unknown.xlsx", data_only=True)
          try:
            cells = list(workbook["Cost Comparison"].values)
            self.assertTrue(all(row[1] == "Unknown" for row in cells[1:3]))
            self.assertIn("incumbent costs unknown", str(cells))
            self.assertIn("target-only estimate", str(list(workbook["Executive Summary"].values)))
          finally:
            workbook.close()

  def test_explicit_zero_incumbent_is_known_but_percentage_is_undefined(self):
    for row, line in zip(self.normalized["aws_boq"], self.plan["lines"]):
      row["Monthly"] = 0
      line["source"]["monthlyCost"] = 0
    self.review["baseline"] = baseline_for(self.plan, CONTEXT, self.normalized)
    self.approve_all()
    summary = self.compile()
    self.assertTrue(summary["opportunityStatus"]["incumbentBaselineComplete"])
    self.assertTrue(summary["opportunityStatus"]["qualifiedComparison"])
    self.assertEqual(0, summary["scenarioTotalsMonthly"]["aws"])
    self.assertEqual(-146, summary["scenarioSavings"]["conservative"]["amount"])
    self.assertIsNone(summary["scenarioSavings"]["conservative"]["percent"])
    self.assertIn("zero incumbent baseline", _render_executive_summary(summary))

  def test_invalid_approval_metadata_rejected(self):
    assumption = self.assumption()
    self.approve(assumption)
    for field, value in (("name", ""), ("evidence", " "), ("confirmedAt", "2026-02-30T12:00:00Z"),
                         ("confirmedAt", "2026-01-02"), ("confirmedAt", "2026-01-02T12:00:00"),
                         ("confirmedAt", "2999-01-02T12:00:00Z")):
      with self.subTest(field=field, value=value):
        plan = copy.deepcopy(self.plan)
        plan["opportunityReview"]["assumptions"][0]["approval"][field] = value
        with self.assertRaises(QuotePlanError):
          validate_quote_plan(plan)
    del assumption["approval"]
    with self.assertRaisesRegex(QuotePlanError, "wording is not approval"):
      validate_quote_plan(self.plan)

  def test_missing_gate_approvals_and_pricing_gaps_block_release(self):
    self.assertFalse(self.status()["releaseReady"])
    self.assertEqual("Submission blocked / draft", self.status()["readiness"])
    self.approve_all()
    self.assertFalse(self.status(pricing_blockers=1)["releaseReady"])
    self.assertFalse(self.status(missing_fields=["RTO"])["releaseReady"])
    self.review["gates"] = self.review["gates"][:-1]
    self.assertEqual("missing", self.status()["approvals"]["customer-release"])
    self.assertFalse(self.status()["releaseReady"])

  def test_traceability_nonpriced_justification(self):
    self.review["requirements"][0]["lineIds"] = []
    self.approve_all()
    self.assertFalse(self.status()["releaseReady"])
    self.assertEqual(1, self.status()["coverage"]["untraced"])
    self.review["requirements"][0]["nonPricedJustification"] = "Eligibility certificate only; no chargeable service"
    self.approve_all()
    self.assertTrue(self.status()["releaseReady"])

  def test_legacy_internal_estimate_and_no_definitive_savings(self):
    plan = copy.deepcopy(self.plan)
    del plan["opportunityReview"]
    summary = self.compile(plan)
    status = summary["opportunityStatus"]
    self.assertEqual("Not reviewed / internal estimate", status["readiness"])
    self.assertFalse(status["releaseReady"])
    self.assertEqual(146, summary["scenarioTotalsMonthly"]["conservative"])
    self.assertEqual("N/A", summary["recommendedScenario"])
    self.assertEqual("Not qualified / provisional comparison only", summary["scenarioSavings"]["conservative"])

  def test_direct_legacy_workbook_keeps_existing_plan_validations(self):
    normalized = copy.deepcopy(self.normalized)
    normalized["aws_boq"][0]["Plan Validations"] = ["Legacy explicit blocker"]
    with patch("build_workbook.resolve_best_price", return_value=PRICE):
      summary = build_workbook(normalized, CONTEXT, self.root / "legacy-direct.xlsx", "conservative")
    self.assertEqual(1, summary["validateCount"])
    self.assertEqual(73, summary["scenarioPricedSubtotalsMonthly"]["conservative"]["amount"])

  def test_quick_triage_no_compiler_or_pricing_and_empty_lines(self):
    self.plan["targetProvider"] = "gcp"
    self.plan["lines"] = []
    self.review["mode"] = "quick-triage"
    self.review["requirements"][0].update(lineIds=[], nonPricedJustification="Triage qualification only")
    normalized = dict(self.normalized, aws_boq=[])
    self.review["baseline"] = baseline_for(self.plan, CONTEXT, normalized)
    normalized["quote_plan"] = self.plan
    validate_quote_plan(self.plan)
    with patch("build_workbook.resolve_best_price") as resolve, patch("providers.targets.require_target_provider") as target:
      summary = build_workbook(normalized, CONTEXT, self.root / "triage.xlsx", "conservative")
    resolve.assert_not_called()
    target.assert_not_called()
    self.assertEqual("Quick triage / not priced", summary["opportunityStatus"]["readiness"])
    self.assertEqual({}, summary["scenarioTotalsMonthly"])
    self.assertFalse(summary["opportunityStatus"]["releaseReady"])

  def test_aws_gcp_working_estimate_compilation_is_not_claimed(self):
    for provider in ("aws", "gcp"):
      with self.subTest(provider=provider):
        plan = copy.deepcopy(self.plan)
        plan["targetProvider"] = provider
        for line in plan["lines"]:
          line["target"]["provider"] = provider
        with self.assertRaisesRegex(ValueError, "not 'compile'"):
          self.compile(plan)

  def test_status_surfaces_match_workbook_summary_and_executive(self):
    self.assumption()
    summary = self.compile()
    status = summary["opportunityStatus"]
    workbook = load_workbook(self.root / "quote.xlsx", data_only=True)
    self.addCleanup(workbook.close)
    values = dict(list(workbook["Opportunity Review"].values)[1:len(status) + 1])
    for key, value in status.items():
      self.assertEqual(value, values[key] if isinstance(value, str) else json.loads(values[key]))
    executive = _render_executive_summary(summary)
    self.assertIn(status["readiness"], executive)
    self.assertIn(status["eligibility"], executive)
    self.assertIn(status["nextAction"], executive)
    self.assertIn(status["owner"], executive)
    self.assertIn("Material assumption A1", executive)
    self.assertEqual("N/A", summary["recommendedScenario"])

  def run_pipeline(self, plan, name):
    plan_path = self.root / "input-plan.json"
    plan_path.write_text(json.dumps(plan), encoding="utf-8")
    normalized_path = self.root / "input-normalized.json"
    normalized_path.write_text(json.dumps(self.normalized), encoding="utf-8")
    specs, boq = self.specs_path, self.boq_path
    argv = ["run_pipeline.py", "--specs", str(specs), "--aws-boq", str(boq),
            "--output", str(self.root / name), "--plan", str(plan_path),
            "--normalized-input", str(normalized_path), "--scenario", "conservative"]
    stdout = io.StringIO()
    with patch.object(sys, "argv", argv), patch.dict(os.environ), redirect_stdout(stdout), \
         patch("build_workbook.resolve_best_price", return_value=PRICE), patch("pricing.flush_cache"):
      pipeline_main()
    return json.loads(stdout.getvalue())

  def test_fresh_reused_and_changed_baseline_status_consistency(self):
    self.assumption()
    self.approve_all()
    first = self.run_pipeline(self.plan, "pipeline.xlsx")
    second = self.run_pipeline(self.plan, "pipeline.xlsx")
    self.assertFalse(first["reused"])
    self.assertTrue(second["reused"])
    self.assertEqual(first["opportunityStatus"], second["opportunityStatus"])
    self.assertTrue(second["opportunityStatus"]["releaseReady"])
    self.plan["lines"][0]["source"]["quantity"] = 999
    third = self.run_pipeline(self.plan, "pipeline.xlsx")
    self.assertFalse(third["reused"])
    self.assertFalse(third["opportunityStatus"]["releaseReady"])
    self.assertEqual(2, len(third["opportunityStatus"]["blockedLineIds"]))
    fourth = self.run_pipeline(self.plan, "pipeline.xlsx")
    self.assertTrue(fourth["reused"])
    self.assertEqual(third["opportunityStatus"], fourth["opportunityStatus"])

  def test_changed_original_sources_block_supplied_normalized_costs_and_cached_release(self):
    self.approve_all()
    first = self.run_pipeline(self.plan, "source-binding.xlsx")
    self.assertTrue(first["opportunityStatus"]["releaseReady"])
    for which in ("specs", "boq", "both"):
      with self.subTest(which=which):
        self.specs_path.write_text('{"changed":true}' if which in ("specs", "both") else "{}", encoding="utf-8")
        self.boq_path.write_text("Service,Quantity\nEC2,999\n" if which in ("boq", "both") else "Service,Quantity\nEC2,730\n", encoding="utf-8")
        changed = self.run_pipeline(self.plan, "source-binding.xlsx")
        self.assertFalse(changed["reused"])
        self.assertFalse(changed["opportunityStatus"]["baselineCurrent"])
        self.assertFalse(changed["opportunityStatus"]["releaseReady"])
        self.assertEqual(2, len(changed["opportunityStatus"]["blockedLineIds"]))
        self.assertTrue(changed["opportunityStatus"]["sourceBindingIssues"])
        summary = json.loads(Path(changed["summary"]).read_text(encoding="utf-8"))
        self.assertEqual(2, summary["validateCount"])
        self.assertEqual("N/A ([VALIDATE])", summary["scenarioTotalsMonthly"]["conservative"])
        reused = self.run_pipeline(self.plan, "source-binding.xlsx")
        self.assertTrue(reused["reused"])
        self.assertEqual(changed["opportunityStatus"], reused["opportunityStatus"])

  def test_unbound_original_inputs_block_review_but_keep_legacy_internal_estimates(self):
    self.normalized.pop("sourceBindings")
    self.review["baseline"] = baseline_for(self.plan, CONTEXT, self.normalized)
    self.approve_all()
    reviewed = self.run_pipeline(self.plan, "unbound.xlsx")
    self.assertFalse(reviewed["opportunityStatus"]["baselineCurrent"])
    self.assertEqual(2, len(reviewed["opportunityStatus"]["blockedLineIds"]))
    legacy = copy.deepcopy(self.plan)
    legacy.pop("opportunityReview")
    self.specs_path.write_text('{"new":true}', encoding="utf-8")
    internal = self.run_pipeline(legacy, "legacy-unbound.xlsx")
    self.assertEqual("Not reviewed / internal estimate", internal["opportunityStatus"]["readiness"])
    summary = json.loads(Path(internal["summary"]).read_text(encoding="utf-8"))
    self.assertEqual(146, summary["scenarioTotalsMonthly"]["conservative"])

  def test_rebind_cannot_bless_normalized_data_from_changed_sources(self):
    self.approve_all()
    self.specs_path.write_text('{"changed":true}', encoding="utf-8")
    plan_path, normalized_path = self.root / "input.json", self.root / "normalized.json"
    plan_path.write_text(json.dumps(self.plan), encoding="utf-8")
    normalized_path.write_text(json.dumps(self.normalized), encoding="utf-8")
    output = self.root / "rebound.json"
    result = subprocess.run(
      [sys.executable, str(ROOT / "scripts" / "opportunity_review.py"), "rebind",
       "--plan", str(plan_path), "--normalized", str(normalized_path), "--output", str(output),
       "--scenario", "conservative"], capture_output=True, text=True,
    )
    self.assertNotEqual(0, result.returncode)
    self.assertIn("rerun parse_inputs.py", result.stderr)
    self.assertFalse(output.exists())

  def test_schema_runtime_vocabulary_and_invalid_numeric_values(self):
    allowed = {"$schema", "title", "type", "required", "additionalProperties", "properties", "$defs",
               "$ref", "const", "enum", "minLength", "pattern", "items", "uniqueItems",
               "format", "minimum", "maximum", "exclusiveMinimum"}
    def walk(schema):
      self.assertFalse(set(schema) - allowed)
      for key in ("properties", "$defs"):
        for sub in schema.get(key, {}).values():
          walk(sub)
      if "items" in schema:
        walk(schema["items"])
    walk(SCHEMA)
    for value in (True, float("nan"), float("inf")):
      with self.subTest(value=value):
        plan = copy.deepcopy(self.plan)
        plan["lines"][0]["source"]["quantity"] = value
        with self.assertRaises(QuotePlanError):
          validate_quote_plan(plan)
        plan = copy.deepcopy(self.plan)
        plan["opportunityReview"]["requirements"][0].update(priority="scored", score=value, weight=1)
        with self.assertRaises(QuotePlanError):
          validate_quote_plan(plan)

  def test_template_rebind_and_explicit_approval_cli(self):
    legacy = copy.deepcopy(self.plan)
    legacy.pop("opportunityReview")
    input_path = self.root / "plan.json"
    input_path.write_text(json.dumps(legacy), encoding="utf-8")
    normalized = self.root / "normalized.json"
    normalized.write_text(json.dumps(self.normalized), encoding="utf-8")
    output = self.root / "reviewed.json"
    base = [sys.executable, str(ROOT / "scripts" / "opportunity_review.py")]
    shared = ["--normalized", str(normalized), "--scenario", "conservative"]
    result = subprocess.run(base + ["template", "--plan", str(input_path), "--output", str(output)] + shared,
                            capture_output=True, text=True)
    self.assertEqual(0, result.returncode, result.stderr)
    created = json.loads(output.read_text(encoding="utf-8"))
    self.assertFalse(created["opportunityReview"]["baselineComplete"])
    self.assertEqual(list(GATES), [g["id"] for g in created["opportunityReview"]["gates"]])
    result = subprocess.run(base + ["approve-gate", "--plan", str(output), "--output", str(output),
                                   "--id", "requirements", "--name", "Reviewer", "--confirmed-at",
                                   "2026-01-02T12:00:00Z", "--evidence", "User message 123"] + shared,
                            capture_output=True, text=True)
    self.assertEqual(0, result.returncode, result.stderr)
    created = json.loads(output.read_text(encoding="utf-8"))
    self.assertEqual("approved", created["opportunityReview"]["gates"][0]["status"])
    result = subprocess.run(base + ["rebind", "--plan", str(output), "--output", str(output)] + shared,
                            capture_output=True, text=True)
    self.assertEqual(0, result.returncode, result.stderr)
    created = json.loads(output.read_text(encoding="utf-8"))
    self.assertFalse(created["opportunityReview"]["baselineComplete"])
    self.assertTrue(all(g["status"] == "pending" and "approval" not in g for g in created["opportunityReview"]["gates"]))


if __name__ == "__main__":
  unittest.main()
