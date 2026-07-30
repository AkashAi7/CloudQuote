import argparse
from concurrent.futures import ThreadPoolExecutor
import json
import re
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from pricing import fetch_prices, resolve_best_price, select_best_price
from providers import resolve_catalog_price
import yaml


SCENARIOS = ["conservative", "moderate", "aggressive"]
AWS_SERVICE_ALIASES = {
  "amazon ec2": "EC2",
  "ec2": "EC2",
  "amazon ebs": "EBS",
  "ebs": "EBS",
  "amazon s3": "S3",
  "s3": "S3",
  "amazon rds": "RDS",
  "rds": "RDS",
  "amazon dynamodb": "DynamoDB",
  "dynamodb": "DynamoDB",
  "aws lambda": "Lambda",
  "lambda": "Lambda",
  "amazon eks": "EKS",
  "eks": "EKS",
  "elb": "ELB",
  "alb": "ALB",
  "application load balancer": "ALB",
  "network load balancer": "ELB",
  "redshift": "Redshift",
  "amazon redshift": "Redshift",
  "cloudfront": "CloudFront",
  "amazon cloudfront": "CloudFront",
  "vpc": "VPC",
  "amazon vpc": "VPC",
  "github": "GitHub",
  "github.com": "GitHub",
}

# Target services that publish reserved-capacity meters in the Azure Retail Prices API.
# Lines mapped to these services attempt Reservation pricing for 1-year/3-year scenarios and
# fall back to Consumption with an explicit [RESERVATION_FALLBACK] note when no term price exists.
RESERVATION_CAPABLE_TARGET_SERVICES = {
  "virtual machines",
  "azure database for mysql",
  "azure database for postgresql",
  "azure database for mariadb",
  "sql database",
  "azure cosmos db",
  "redis cache",
  "azure cache for redis",
  "azure synapse analytics",
  "sql data warehouse",
  "azure data explorer",
}


def _num(v: Any) -> float:
  if isinstance(v, (int, float)):
    return float(v)
  if isinstance(v, str) and v.strip() != "":
    try:
      return float(v)
    except ValueError:
      return 0.0
  return 0.0


def _normalize_aws_service(raw: str) -> str:
  key = raw.strip().lower()
  return AWS_SERVICE_ALIASES.get(key, raw.strip())


def _scenario_price_query(scenario: str) -> Dict[str, str | None]:
  if scenario == "conservative":
    return {"price_type": "Consumption", "reservation_term": None}
  if scenario == "moderate":
    return {"price_type": "Reservation", "reservation_term": "1 Year"}
  return {"price_type": "Reservation", "reservation_term": "3 Years"}


def _reservation_months(reservation_term: str | None) -> int:
  if reservation_term == "1 Year":
    return 12
  if reservation_term == "3 Years":
    return 36
  return 1


def _uom_multiplier(unit_of_measure: str) -> float:
  """Return the pack size encoded in a retail unitOfMeasure such as '100 Hours' or '1M Operations'."""
  match = re.match(r"\s*(\d+(?:\.\d+)?)\s*([km])?\b", (unit_of_measure or "").lower())
  if not match:
    return 1.0
  value = float(match.group(1))
  suffix = match.group(2)
  if suffix == "k":
    value *= 1_000.0
  elif suffix == "m":
    value *= 1_000_000.0
  return value if value > 0 else 1.0


def _monthly_from_unit(unit_price: float, unit_of_measure: str, qty: float) -> tuple[float | None, str]:
  uom = (unit_of_measure or "").lower()
  if uom == "10":
    return unit_price * qty, "per_10_units"
  pack = _uom_multiplier(uom)
  pack_note = "" if pack == 1.0 else f"_per_{int(pack) if pack.is_integer() else pack}_units"
  if "gib/hour" in uom or "gb/hour" in uom:
    return unit_price * qty * 730.0 / pack, f"capacity_hourly_to_monthly_730h{pack_note}"
  if "hour" in uom:
    return unit_price * qty * 730.0 / pack, f"hourly_to_monthly_730h{pack_note}"
  if "day" in uom:
    return unit_price * qty * 30.4167 / pack, f"daily_to_monthly_30_4167d{pack_note}"
  if "month" in uom:
    return unit_price * qty / pack, f"monthly{pack_note}"
  if pack != 1.0:
    return unit_price * qty / pack, f"per_unit{pack_note}"
  # A bare consumption unit such as '1 GB', '1' or '1 Rotation' carries no time dimension, so the
  # meter is charged per unit consumed. The BOQ contract states quantities are monthly volumes,
  # which makes the direct product the monthly cost. The note keeps that assumption reviewable.
  return unit_price * qty, "per_unit_monthly_quantity_assumed"


def _extract_capacity_number(text: str, unit: str) -> float | None:
  match = re.search(r"(\d+(?:\.\d+)?)\s*" + re.escape(unit), text.lower())
  if not match:
    return None
  return float(match.group(1))


def _extract_vcpu(text: str) -> float | None:
  match = re.search(r"(\d+(?:\.\d+)?)\s*vcpu", text.lower())
  if not match:
    return None
  return float(match.group(1))

def _extract_requests(text: str) -> float | None:
  # Supports forms like "12M requests" or "12000000 requests"
  million = re.search(r"(\d+(?:\.\d+)?)\s*m\s*requests", text.lower())
  if million:
    return float(million.group(1)) * 1_000_000.0
  plain = re.search(r"(\d+(?:\.\d+)?)\s*requests", text.lower())
  if plain:
    return float(plain.group(1))
  return None

def _extract_rw_units(text: str) -> float | None:
  million = re.search(r"(\d+(?:\.\d+)?)\s*m\s*r/w", text.lower())
  if million:
    return float(million.group(1)) * 1_000_000.0
  plain = re.search(r"(\d+(?:\.\d+)?)\s*r/w", text.lower())
  if plain:
    return float(plain.group(1))
  return None


def _extract_users(text: str) -> float | None:
  match = re.search(r"(\d+(?:\.\d+)?)\s*(?:users?|seats?|licenses?)", text.lower())
  return float(match.group(1)) if match else None

def _billable_quantity(service: str, qty: float, capacity: str) -> float:
  lower = capacity.lower()
  if service == "EBS":
    tb = _extract_capacity_number(lower, "tb")
    gb = _extract_capacity_number(lower, "gb")
    if tb is not None:
      return tb * 1024.0
    if gb is not None:
      return gb
  if service == "S3":
    tb = _extract_capacity_number(lower, "tb")
    gb = _extract_capacity_number(lower, "gb")
    if tb is not None:
      return tb * 1024.0
    if gb is not None:
      return gb
  if service == "CloudFront":
    tb = _extract_capacity_number(lower, "tb")
    gb = _extract_capacity_number(lower, "gb")
    if tb is not None:
      return tb * 1024.0
    if gb is not None:
      return gb
  if service == "Lambda":
    req = _extract_requests(lower)
    if req is not None:
      # Azure Functions sample meter is per-10 executions.
      return req / 10.0
  if service == "RDS":
    vcpu = _extract_vcpu(lower)
    if vcpu is not None:
      return vcpu * qty
  if service == "DynamoDB":
    rw = _extract_rw_units(lower)
    if rw is not None:
      # Cosmos autoscale meter is per 100 RU/s-hour equivalent in this sample mapping.
      return max(rw / 100.0, 1.0)
  if service == "GitHub":
    users = _extract_users(lower)
    return users if users is not None else qty
  if service == "EKS":
    nodes = re.search(r"(\d+(?:\.\d+)?)\s*nodes", lower)
    if nodes:
      return float(nodes.group(1))
    return qty
  if service == "VPC":
    return qty
  if service == "Redshift":
    return qty
  return qty


def _instance_map() -> Dict[str, Any]:
  mapping_path = Path(__file__).resolve().parent.parent / "mappings" / "instance_sizes.yaml"
  if not mapping_path.exists():
    return {}
  data = yaml.safe_load(mapping_path.read_text(encoding="utf-8")) or {}
  return data.get("instances", {})


def _candidate_skus(service: str, source_sku: str, current_azure_sku: str) -> List[str]:
  if service != "EC2":
    return [current_azure_sku] if current_azure_sku else []
  instance_map = _instance_map()
  candidates = [current_azure_sku] if current_azure_sku else []
  mapped = instance_map.get(source_sku, {}).get("azure_candidates", [])
  for candidate in mapped:
    if candidate not in candidates:
      candidates.append(candidate)
  return candidates


def _price_profile(service: str, azure_service: str, azure_sku: str, capacity: str, source_sku: str) -> Dict[str, Any]:
  if service == "EC2":
    return {
      "service_name": "Virtual Machines",
      "sku_name": azure_sku,
      "sku_candidates": _candidate_skus(service, source_sku, azure_sku),
      "meter_contains": None,
      "product_contains": None,
      "reservation_supported": True,
      "strict_sku": True,
      "capacity_note": "",
    }
  if service == "EBS":
    return {
      "service_name": "Storage",
      "sku_name": None,
      "sku_candidates": [],
      "meter_contains": "Provisioned Capacity",
      "product_contains": "Premium SSD v2",
      "reservation_supported": False,
      "capacity_note": "IOPS and throughput surcharges may require validation.",
      "conversion_validation": "Managed disk pricing is incomplete without tier, provisioned IOPS, and throughput assumptions",
    }
  if service == "S3":
    return {
      "service_name": "Storage",
      "sku_name": "Hot LRS",
      "sku_candidates": ["Hot LRS"],
      "meter_contains": "Data Stored",
      "product_contains": "Block Blob",
      "reservation_supported": False,
      "capacity_note": "Transaction and redundancy surcharges may require validation.",
      "conversion_validation": "Blob storage pricing is incomplete without redundancy, transaction, retrieval, and egress assumptions",
    }
  if service == "ALB" or "application" in source_sku.lower() or azure_service == "Azure Application Gateway":
    return {
      "service_name": "Application Gateway",
      "sku_name": None,
      "sku_candidates": [],
      "meter_contains": "Fixed Cost",
      "product_contains": "Application Gateway Standard v2",
      "reservation_supported": False,
      "capacity_note": "Capacity unit charges may require validation.",
      "conversion_validation": "Application Gateway pricing requires gateway hours, capacity units, and processed data",
    }
  if service == "ELB":
    return {
      "service_name": "Load Balancer",
      "sku_name": None,
      "sku_candidates": [],
      "meter_contains": None,
      "product_contains": None,
      "reservation_supported": False,
      "capacity_note": "Data processed and rule dimensions may require validation.",
      "conversion_validation": "Load Balancer pricing requires rule hours and processed-data dimensions",
    }
  if service == "RDS":
    return {
      "service_name": "Azure Database for PostgreSQL",
      "sku_name": None,
      "sku_candidates": [],
      "meter_contains": "vCore",
      "product_contains": "Flexible Server",
      "reservation_supported": False,
      "capacity_note": "Storage, backup, and HA overhead may require validation.",
      "conversion_validation": "Database pricing is incomplete without storage, backup, high-availability, and compute-hour assumptions",
    }
  if service == "Lambda":
    return {
      "service_name": "Functions",
      "sku_name": "On Demand",
      "sku_candidates": ["On Demand"],
      "meter_contains": "Total Executions",
      "product_contains": "Consumption",
      "reservation_supported": False,
      "capacity_note": "Execution duration and memory charges may require validation.",
      "conversion_validation": "Function execution count alone is insufficient without execution duration and memory consumption",
    }
  if service == "DynamoDB":
    return {
      "service_name": "Azure Cosmos DB",
      "sku_name": "AP4",
      "sku_candidates": ["AP4"],
      "meter_contains": "100 RUs",
      "product_contains": "autoscale",
      "reservation_supported": False,
      "capacity_note": "Request unit mapping requires validation.",
      "conversion_validation": "DynamoDB operations cannot be converted to provisioned Cosmos RU/s without RU-per-operation and workload-distribution assumptions",
    }
  if service == "CloudFront":
    return {
      "service_name": "Azure Front Door Service",
      "sku_name": None,
      "sku_candidates": [],
      "meter_contains": None,
      "product_contains": None,
      "reservation_supported": False,
      "capacity_note": "Front Door meter mapping requires validation.",
      "force_validate": True,
    }
  if service == "EKS":
    return {
        "service_name": "Azure Kubernetes Service",
      "sku_name": None,
      "sku_candidates": [],
        "meter_contains": None,
        "product_contains": "Managed Cluster",
      "reservation_supported": False,
      "capacity_note": "AKS worker node VM costs should be priced separately from cluster management.",
      "conversion_validation": "AKS management pricing excludes required worker-node compute and attached infrastructure",
    }
  if service == "Redshift":
    return {
      "service_name": "Azure Synapse Analytics",
      "sku_name": None,
      "sku_candidates": [],
      "meter_contains": "DWU",
      "product_contains": "Dedicated SQL Pool",
      "reservation_supported": False,
      "capacity_note": "Synapse storage and query workload variability may require validation.",
      "conversion_validation": "Warehouse pricing requires a defensible source-capacity to DWU mapping plus storage assumptions",
    }
  if service == "VPC":
    return {
      "service_name": "Virtual Network",
      "sku_name": None,
      "sku_candidates": [],
      "meter_contains": None,
      "product_contains": None,
      "reservation_supported": False,
      "capacity_note": "VNet pricing is componentized (NAT, peering, egress, private endpoints) and requires validation.",
      "force_validate": True,
    }
  if service == "GitHub":
    return {
      "service_name": "GitHub",
      "sku_name": azure_sku or source_sku,
      "sku_candidates": [],
      "meter_contains": None,
      "product_contains": None,
      "reservation_supported": False,
      "capacity_note": "Official GitHub public catalog pricing; add-ons require separate lines.",
      "pricing_provider": "official_catalog",
      "catalog_name": "github",
      "catalog_sku": azure_sku or source_sku,
    }

  return {
    "service_name": azure_service,
    "sku_name": azure_sku or None,
    "sku_candidates": [azure_sku] if azure_sku else [],
    "meter_contains": None,
    "product_contains": None,
    "reservation_supported": azure_service.strip().lower() in RESERVATION_CAPABLE_TARGET_SERVICES,
    "capacity_note": "",
    "force_validate": False,
  }


def _lookup_live_price(
  region: str,
  currency: str,
  profile: Dict[str, Any],
  scenario: str,
  billable_qty: float,
  source_unit: str = "",
  source_sku: str = "",
  capacity: str = "",
) -> Dict[str, Any]:
  if profile.get("pricing_provider") == "official_catalog":
    return resolve_catalog_price(
      str(profile["catalog_name"]),
      str(profile.get("catalog_sku", source_sku)),
      billable_qty,
      currency,
    )
  if profile.get("force_validate"):
    return {
      "unit_price": 0.0,
      "monthly": 0.0,
      "annual": 0.0,
      "uom": "",
      "price_date": "",
      "note": "[VALIDATE] Componentized pricing required",
      "validate": True,
    }
  # Zone-priced global services (CDN, DNS, Front Door) publish no meters under a physical region,
  # so a line may name the pricing region its meter actually lives in.
  region = str(profile.get("price_region") or region)
  conversion_validation = str(profile.get("conversion_validation", "")).strip()
  if conversion_validation:
    intended_meter = profile.get("meter_contains") or profile.get("product_contains") or "unspecified"
    return {
      "unit_price": 0.0,
      "monthly": 0.0,
      "annual": 0.0,
      "uom": "",
      "price_date": "",
      "note": f"[VALIDATE] {conversion_validation}; intended meter filter: {intended_meter}",
      "validate": True,
      "resolved_sku": profile.get("sku_name") or source_sku,
      "source": "NONE",
      "source_links": [],
      "source_note": "Pricing lookup skipped until required conversion assumptions are supplied",
    }

  query = _scenario_price_query(scenario)
  if not profile["reservation_supported"]:
    query = {"price_type": "Consumption", "reservation_term": None}

  chosen_sku = profile["sku_name"]
  sku_candidates = profile.get("sku_candidates", []) or ([profile["sku_name"]] if profile["sku_name"] else [])
  resolved = None
  for candidate in sku_candidates:
    resolved = resolve_best_price(
      region=region,
      currency_code=currency,
      service_name=profile["service_name"],
      sku_name=candidate,
      meter_name=None,
      price_type=str(query["price_type"]),
      reservation_term=query["reservation_term"],
      target_sku=candidate,
      meter_contains=profile["meter_contains"],
      product_contains=profile["product_contains"],
      select_service_name=profile["service_name"],
      search_hint=f"{source_sku} {capacity}".strip(),
      strict_contains=bool(profile.get("strict_contains")),
      reviewed_evidence=profile.get("pricing_evidence", []),
    )
    if resolved and resolved.get("best"):
      chosen_sku = candidate
      break

  if (not resolved or not resolved.get("best")) and not profile.get("strict_sku") and profile["service_name"]:
    resolved = resolve_best_price(
      region=region,
      currency_code=currency,
      service_name=profile["service_name"],
      sku_name=None,
      meter_name=None,
      price_type=str(query["price_type"]),
      reservation_term=query["reservation_term"],
      target_sku=chosen_sku,
      meter_contains=profile["meter_contains"],
      product_contains=profile["product_contains"],
      select_service_name=profile["service_name"],
      search_hint=f"{source_sku} {capacity}".strip(),
      strict_contains=bool(profile.get("strict_contains")),
      reviewed_evidence=profile.get("pricing_evidence", []),
    )

  best = resolved.get("best") if resolved else None
  reservation_fallback_note = ""
  reservation_fallback_attempted = False
  if not best and str(query["price_type"]) == "Reservation":
    reservation_fallback_attempted = True
    # If reservation term is missing from sources, fallback to consumption so scenario still resolves with explicit note.
    resolved = resolve_best_price(
      region=region,
      currency_code=currency,
      service_name=profile["service_name"],
      sku_name=chosen_sku,
      meter_name=None,
      price_type="Consumption",
      reservation_term=None,
      target_sku=chosen_sku,
      meter_contains=profile["meter_contains"],
      product_contains=profile["product_contains"],
      select_service_name=profile["service_name"],
      search_hint=f"{source_sku} {capacity}".strip(),
      strict_contains=bool(profile.get("strict_contains")),
      reviewed_evidence=profile.get("pricing_evidence", []),
    )
    best = resolved.get("best") if resolved else None
    if best:
      reservation_fallback_note = "[RESERVATION_FALLBACK] consumption price used"

  if not best:
    resolved_links = list(resolved.get("links", [])) if resolved else []
    resolved_note = str(resolved.get("note", "")) if resolved else ""
    return {
      "unit_price": 0.0,
      "monthly": 0.0,
      "annual": 0.0,
      "uom": "",
      "price_date": "",
      "note": "[VALIDATE] SKU/product mapping unresolved after API+WEB lookup",
      "validate": True,
      "source": "NONE",
      "source_links": resolved_links,
      "source_note": resolved_note,
      "fallback_attempted": bool(resolved and resolved.get("fallbackAttempted")) or reservation_fallback_attempted,
      "fallback_succeeded": False,
      "fallback_latency_ms": float(resolved.get("fallbackLatencyMs", 0.0)) if resolved else 0.0,
    }

  retail_price = float(best.get("retailPrice", 0.0))
  unit_of_measure = str(best.get("unitOfMeasure", ""))
  is_web_evidence = bool(
    resolved
    and (
      resolved.get("source") == "WEB"
      or resolved.get("originSource") == "WEB"
      or best.get("productName") == "Web search estimate"
      or unit_of_measure == "Web Estimated Unit"
    )
  )
  if is_web_evidence:
    return {
      "unit_price": retail_price,
      "monthly": 0.0,
      "annual": 0.0,
      "uom": unit_of_measure,
      "price_date": str(best.get("effectiveStartDate", "")),
      "note": "[VALIDATE] Web estimate has no verified Azure meter dimension",
      "validate": True,
      "resolved_sku": chosen_sku or str(best.get("skuName", "")),
      "source": str(resolved.get("source", "WEB")),
      "source_links": list(resolved.get("links", [])),
      "source_note": str(resolved.get("note", "")),
    }
  if str(query["price_type"]) == "Reservation":
    months = _reservation_months(query["reservation_term"])
    monthly = (retail_price / months) * billable_qty
    conversion_note = f"reservation_upfront_amortized_{months}m"
  else:
    source_unit_lower = (source_unit or "").lower()
    if "hour" in source_unit_lower and "hour" in unit_of_measure.lower():
      monthly = retail_price * billable_qty
      conversion_note = "source_quantity_already_monthly_hours"
    elif "month" in source_unit_lower and "month" in unit_of_measure.lower():
      monthly = retail_price * billable_qty
      conversion_note = "source_quantity_monthly"
    else:
      monthly, conversion_note = _monthly_from_unit(retail_price, unit_of_measure, billable_qty)

  if monthly is None:
    return {
      "unit_price": retail_price,
      "monthly": 0.0,
      "annual": 0.0,
      "uom": unit_of_measure,
      "price_date": str(best.get("effectiveStartDate", "")),
      "note": conversion_note,
      "validate": True,
    }

  return {
    "unit_price": retail_price,
    "monthly": monthly,
    "annual": monthly * 12.0,
    "uom": unit_of_measure,
    "price_date": str(best.get("effectiveStartDate", "")),
    "note": f"{conversion_note}; {reservation_fallback_note}" if reservation_fallback_note else conversion_note,
    "validate": False,
    "resolved_sku": chosen_sku or str(best.get("skuName", "")),
    "source": str(resolved.get("source", "API")) if resolved else "API",
    "source_links": list(resolved.get("links", [])) if resolved else [],
    "source_note": str(resolved.get("note", "")) if resolved else "",
    "fallback_attempted": bool(resolved and resolved.get("fallbackAttempted")) or reservation_fallback_attempted,
    "fallback_succeeded": bool(resolved and resolved.get("fallbackSucceeded")) or bool(reservation_fallback_note),
    "fallback_latency_ms": float(resolved.get("fallbackLatencyMs", 0.0)) if resolved else 0.0,
  }


def _price_cache_key(
  region: str,
  currency: str,
  scenario: str,
  billable_qty: float,
  source_unit: str,
  profile: Dict[str, Any],
) -> str:
  return json.dumps(
    {
      "region": region,
      "currency": currency,
      "scenario": scenario,
      "billable_qty": round(float(billable_qty), 6),
      "source_unit": source_unit.lower(),
      "service": profile.get("service_name", ""),
      "sku": profile.get("sku_name", ""),
      "sku_candidates": profile.get("sku_candidates", []),
      "meter_contains": profile.get("meter_contains", ""),
      "product_contains": profile.get("product_contains", ""),
      "price_region": profile.get("price_region", ""),
      "strict_contains": bool(profile.get("strict_contains", False)),
      "reservation_supported": bool(profile.get("reservation_supported", False)),
      "force_validate": bool(profile.get("force_validate", False)),
      "conversion_validation": profile.get("conversion_validation", ""),
      "pricing_provider": profile.get("pricing_provider", ""),
      "catalog_name": profile.get("catalog_name", ""),
      "catalog_sku": profile.get("catalog_sku", ""),
      "pricing_evidence": profile.get("pricing_evidence", []),
    },
    sort_keys=True,
  )


_TRUTHY = {"1", "true", "yes", "y"}


def _apply_plan_overrides(profile: Dict[str, Any], row: Dict[str, Any]) -> Dict[str, Any]:
  """Let a reviewed quote plan or a derived BOQ pin the meter a line must resolve against.

  Two override channels are supported:

  * Documented assumptions on a line mean the analyst has supplied the dimensions the built-in
    guardrail was waiting for, so the blocker is replaced by whatever residual validations the
    plan still declares. Lines without assumptions or validations keep the built-in guardrail.
  * Optional ``Meter Contains`` / ``Product Contains`` / ``Force Validate`` columns on the source
    BOQ pin the retail meter directly. This is required for RFP-derived inputs, where one target
    service publishes many meters (deployment vs capacity unit vs data processed) and only the
    tender author knows which one the line is asking for.
  """
  overridden = dict(profile)
  changed = False

  plan_assumptions = [str(item) for item in row.get("Mapping Assumptions", []) if str(item).strip()]
  plan_validations = [str(item) for item in row.get("Plan Validations", []) if str(item).strip()]
  if plan_assumptions or plan_validations:
    overridden["conversion_validation"] = "; ".join(plan_validations)
    changed = True

  meter_contains = str(row.get("Meter Contains", "")).strip()
  product_contains = str(row.get("Product Contains", "")).strip()
  if meter_contains:
    overridden["meter_contains"] = meter_contains
    changed = True
  if product_contains:
    overridden["product_contains"] = product_contains
    changed = True
  if meter_contains or product_contains:
    # Pinning a meter is the analyst answering the question the built-in guardrail was asking,
    # so the blocker it raised no longer applies to this line. The pin is also authoritative:
    # if it matches nothing the line must fail loudly rather than resolve to a different meter.
    overridden["conversion_validation"] = ""
    overridden["force_validate"] = False
    overridden["strict_contains"] = True
    changed = True
  price_region = str(row.get("Price Region", "")).strip()
  if price_region:
    overridden["price_region"] = price_region
    changed = True
  if str(row.get("Force Validate", "")).strip().lower() in _TRUTHY:
    overridden["force_validate"] = True
    changed = True
  pricing_evidence = row.get("Pricing Evidence", [])
  if isinstance(pricing_evidence, list) and pricing_evidence:
    overridden["pricing_evidence"] = pricing_evidence
    changed = True

  return overridden if changed else profile


def _prefetch_prices(
  rows: List[Dict[str, Any]],
  scenarios: List[str],
  region: str,
  currency: str,
) -> Dict[str, Dict[str, Any]]:
  requests_by_key: Dict[str, tuple] = {}
  for row in rows:
    service = _normalize_aws_service(str(row.get("Service", row.get("service", ""))))
    sku = str(row.get("Instance/SKU", row.get("instance", row.get("sku", ""))))
    qty = _num(row.get("Quantity", row.get("quantity", 1)))
    source_unit = str(row.get("Source Unit", ""))
    capacity = str(row.get("Capacity", f"Derived from {sku}"))
    azure_service = str(row.get("Azure Service", "Azure mapped service"))
    azure_sku = str(row.get("Azure SKU", sku or "[VALIDATE]"))
    billable_qty = _billable_quantity(service, qty, capacity)
    profile = _price_profile(service, azure_service, azure_sku, capacity, sku)
    profile = _apply_plan_overrides(profile, row)
    for scenario in scenarios:
      key = _price_cache_key(region, currency, scenario, billable_qty, source_unit, profile)
      requests_by_key.setdefault(
        key,
        (region, currency, profile, scenario, billable_qty, source_unit, sku, capacity),
      )

  if not requests_by_key:
    return {}
  with ThreadPoolExecutor(max_workers=min(6, len(requests_by_key))) as executor:
    futures = {
      key: executor.submit(
        _lookup_live_price,
        request[0],
        request[1],
        request[2],
        request[3],
        request[4],
        source_unit=request[5],
        source_sku=request[6],
        capacity=request[7],
      )
      for key, request in requests_by_key.items()
    }
    return {key: future.result() for key, future in futures.items()}


def _line_note(scenario: str, env: str, apply_hybrid_benefit: bool, uses_reservation: bool) -> str:
  notes: List[str] = []
  if uses_reservation:
    notes.append("Reservation")
  if scenario == "aggressive" and env.lower() != "prod":
    notes.append("Dev/Test pricing")
  if apply_hybrid_benefit:
    notes.append("Hybrid Benefit applied")
  return "; ".join(notes)


def _should_apply_hybrid_benefit(service: str, row: Dict[str, Any], specs: Dict[str, Any]) -> bool:
  licenses = [str(item).lower() for item in specs.get("owned_licenses", [])]
  row_text = json.dumps(row).lower()
  if service == "EC2" and "windows" in row_text and any("windows" in item for item in licenses):
    return True
  if service == "RDS" and "sql" in row_text and any("sql server" in item for item in licenses):
    return True
  return False


def _build_pricing_metrics(comparison_rows: List[Dict[str, Any]]) -> Dict[str, Any]:
  source_counts = {"API": 0, "WEB": 0, "CACHE": 0, "NONE": 0}
  for row in comparison_rows:
    source = str(row.get("source", "API"))
    source_counts[source] = source_counts.get(source, 0) + 1
  pricing_line_count = len(comparison_rows)
  unresolved_count = source_counts.get("NONE", 0)
  fallback_attempts = sum(1 for row in comparison_rows if row.get("fallback_attempted"))
  fallback_successes = sum(1 for row in comparison_rows if row.get("fallback_succeeded"))
  fallback_latency_ms = sum(float(row.get("fallback_latency_ms", 0.0)) for row in comparison_rows if row.get("fallback_attempted"))
  return {
    "sourceCounts": source_counts,
    "pricingMetrics": {
      "lineCount": pricing_line_count,
      "unresolvedCount": unresolved_count,
      "unresolvedRate": round(unresolved_count / pricing_line_count, 4) if pricing_line_count else 0.0,
      "sourceCoverage": {
        source: {
          "count": count,
          "rate": round(count / pricing_line_count, 4) if pricing_line_count else 0.0,
        }
        for source, count in sorted(source_counts.items())
      },
      "fallbackAttempts": fallback_attempts,
      "fallbackSuccesses": fallback_successes,
      "fallbackSuccessRate": round(fallback_successes / fallback_attempts, 4) if fallback_attempts else 0.0,
      "fallbackLatencyMsTotal": round(fallback_latency_ms, 3),
      "fallbackLatencyMsAverage": round(fallback_latency_ms / fallback_attempts, 3) if fallback_attempts else 0.0,
    },
  }


def build_workbook(normalized: Dict[str, Any], pricing_meta: Dict[str, Any], output_path: Path, scenario: str = "compare-all") -> Dict[str, Any]:
  wb = Workbook()

  ws1 = wb.active
  ws1.title = "Mapping & Azure BOQ"
  ws1.append([
    "AWS item",
    "Capacity",
    "Azure service",
    "SKU",
    "Region",
    "Qty",
    "Term",
    "Unit price",
    "Monthly",
    "Annual",
    "Price source",
    "Evidence links",
    "Notes/[VALIDATE]",
  ])

  aws_rows = normalized.get("aws_boq", [])
  specs = normalized.get("specs", {})
  region = pricing_meta.get("region", "eastus")
  pricing_date = pricing_meta.get("pricingDate", "")
  currency = pricing_meta.get("currency", "USD")

  target_scenarios = SCENARIOS if scenario == "compare-all" else [scenario]

  comparison_rows: List[Dict[str, Any]] = []
  validation_rows: List[Dict[str, Any]] = []
  price_lookup_cache = _prefetch_prices(aws_rows, target_scenarios, region, currency)
  validate_count = 0
  for row in aws_rows:
    service = _normalize_aws_service(str(row.get("Service", row.get("service", ""))))
    sku = str(row.get("Instance/SKU", row.get("instance", row.get("sku", ""))))
    qty = _num(row.get("Quantity", row.get("quantity", 1)))
    aws_monthly = _num(row.get("Monthly", row.get("monthly", row.get("Monthly Cost", 0))))
    env = str(row.get("Environment", row.get("environment", "prod")))
    source_unit = str(row.get("Source Unit", ""))
    local_unit_price = _num(row.get("Local Azure Unit Price", 0))
    local_total_cost = _num(row.get("Local Azure Total Cost", 0))
    apply_hybrid_benefit = _should_apply_hybrid_benefit(service, row, specs)

    capacity = str(row.get("Capacity", f"Derived from {sku}"))
    azure_service = str(row.get("Azure Service", "Azure mapped service"))
    azure_sku = str(row.get("Azure SKU", sku or "[VALIDATE]"))
    billable_qty = _billable_quantity(service, qty, capacity)
    profile = _price_profile(service, azure_service, azure_sku, capacity, sku)
    profile = _apply_plan_overrides(profile, row)

    for s in target_scenarios:
      cache_key = _price_cache_key(region, currency, s, billable_qty, source_unit, profile)
      live = price_lookup_cache.get(cache_key)
      if live is None:
        live = _lookup_live_price(
          region,
          currency,
          profile,
          s,
          billable_qty,
          source_unit=source_unit,
          source_sku=sku,
          capacity=capacity,
        )

        price_lookup_cache[cache_key] = live
      live = dict(live)

      # Final local fallback: use customer's Azure reference row when available.
      if live.get("validate") and (local_unit_price > 0 or local_total_cost > 0):
        prior_fallback_latency = float(live.get("fallback_latency_ms", 0.0))
        local_monthly = local_total_cost if local_total_cost > 0 else local_unit_price * billable_qty
        local_unit = local_unit_price if local_unit_price > 0 else (local_monthly / billable_qty if billable_qty else 0.0)
        local_notes = ["[LOCAL_CACHE] from workbook Azure OPEX reference"]
        if row.get("Sr. No"):
          local_notes.append(f"Row {row.get('Sr. No')}")
        live = {
          "unit_price": local_unit,
          "monthly": local_monthly,
          "annual": local_monthly * 12.0,
          "uom": str(row.get("Source Unit", "")),
          "price_date": pricing_meta.get("pricingDate", ""),
          "note": "local_reference",
          "validate": False,
          "resolved_sku": str(row.get("Local Azure SKU", "")) or (azure_sku or sku),
          "source": "CACHE",
          "source_links": [],
          "source_note": "; ".join(local_notes),
          "fallback_attempted": True,
          "fallback_succeeded": True,
          "fallback_latency_ms": prior_fallback_latency,
        }
      unit_price = float(live["unit_price"])
      azure_monthly = float(live["monthly"])
      annual = float(live["annual"])
      if live["validate"]:
        validate_count += 1
        validation_rows.append(
          {
            "line": f"{service}:{sku}",
            "scenario": s,
            "reason": str(live["note"]),
            "service": profile["service_name"],
            "source_sku": sku,
            "mapped_sku": azure_sku or "",
          }
        )
      if profile["reservation_supported"]:
        term = "Pay-as-you-go" if s == "conservative" else ("1 Year Reserved" if s == "moderate" else "3 Year Reserved")
      else:
        term = "Pay-as-you-go"
      notes = _line_note(s, env, apply_hybrid_benefit, profile["reservation_supported"] and s in {"moderate", "aggressive"})
      notes = f"{notes}; {live['note']}" if notes else str(live["note"])
      if live.get("source_note"):
        notes = f"{notes}; {live['source_note']}" if notes else str(live["source_note"])
      if profile["capacity_note"]:
        notes = f"{notes}; {profile['capacity_note']}" if notes else profile["capacity_note"]
      mapping_assumptions = [str(item) for item in row.get("Mapping Assumptions", []) if str(item).strip()]
      if mapping_assumptions:
        notes = f"{notes}; Assumptions: {'; '.join(mapping_assumptions)}" if notes else f"Assumptions: {'; '.join(mapping_assumptions)}"
      if live["price_date"] and not pricing_date:
        pricing_date = str(live["price_date"])[:10]

      monthly_cell: float | str = round(azure_monthly, 2)
      annual_cell: float | str = round(annual, 2)
      if live["validate"]:
        monthly_cell = "[VALIDATE]"
        annual_cell = "[VALIDATE]"

      source_links = "\n".join(live.get("source_links", []))
      ws1.append([
        f"{service}:{sku}",
        capacity,
        azure_service,
        str(live.get("resolved_sku", azure_sku or "[VALIDATE]")),
        region,
        billable_qty,
        term,
        round(unit_price, 6),
        monthly_cell,
        annual_cell,
        live.get("source", "API"),
        source_links,
        notes,
      ])
      row_index = ws1.max_row
      source_fill = {
        "API": PatternFill(fill_type="solid", fgColor="C6EFCE"),
        "WEB": PatternFill(fill_type="solid", fgColor="FFF2CC"),
        "CACHE": PatternFill(fill_type="solid", fgColor="FCE4D6"),
        "NONE": PatternFill(fill_type="solid", fgColor="F8CBAD"),
      }.get(str(live.get("source", "API")), PatternFill(fill_type="solid", fgColor="D9E1F2"))
      ws1.cell(row=row_index, column=11).fill = source_fill

      comparison_rows.append(
        {
          "line": f"{service}:{sku}",
          "scenario": s,
          "aws_monthly": aws_monthly,
          "azure_monthly": None if live["validate"] else azure_monthly,
          "source": str(live.get("source", "API")),
          "source_links": list(live.get("source_links", [])),
          "fallback_attempted": bool(live.get("fallback_attempted")),
          "fallback_succeeded": bool(live.get("fallback_succeeded")),
          "fallback_latency_ms": float(live.get("fallback_latency_ms", 0.0)),
        }
      )

  ws2 = wb.create_sheet("Cost Comparison")
  totals = {"aws": 0.0, "conservative": 0.0, "moderate": 0.0, "aggressive": 0.0}
  scenario_has_validate = {"conservative": False, "moderate": False, "aggressive": False}
  # A tender still needs a defensible commercial envelope when one line is pending clarification,
  # so the priced lines are subtotalled separately from the all-or-nothing headline total.
  scenario_unpriced_counts = {"conservative": 0, "moderate": 0, "aggressive": 0}
  line_scenario_values: Dict[str, Dict[str, float | None]] = {}

  def _register_line_value(line: str, scenario_name: str, value: float | None) -> None:
    if line not in line_scenario_values:
      line_scenario_values[line] = {"conservative": None, "moderate": None, "aggressive": None}
    line_scenario_values[line][scenario_name] = value

  if scenario == "compare-all":
    ws2.append([
      "Line",
      "AWS Monthly",
      "Azure conservative",
      "Source conservative",
      "Azure moderate",
      "Source moderate",
      "Azure aggressive",
      "Source aggressive",
      "Web links (if used)",
    ])
    lines = sorted(set(r["line"] for r in comparison_rows))
    for line in lines:
      aws_val = next(r["aws_monthly"] for r in comparison_rows if r["line"] == line)
      cons = next(r["azure_monthly"] for r in comparison_rows if r["line"] == line and r["scenario"] == "conservative")
      mod = next(r["azure_monthly"] for r in comparison_rows if r["line"] == line and r["scenario"] == "moderate")
      agg = next(r["azure_monthly"] for r in comparison_rows if r["line"] == line and r["scenario"] == "aggressive")
      cons_row = next(r for r in comparison_rows if r["line"] == line and r["scenario"] == "conservative")
      mod_row = next(r for r in comparison_rows if r["line"] == line and r["scenario"] == "moderate")
      agg_row = next(r for r in comparison_rows if r["line"] == line and r["scenario"] == "aggressive")
      _register_line_value(line, "conservative", cons)
      _register_line_value(line, "moderate", mod)
      _register_line_value(line, "aggressive", agg)
      cons_cell: float | str = round(cons, 2) if cons is not None else "[VALIDATE]"
      mod_cell: float | str = round(mod, 2) if mod is not None else "[VALIDATE]"
      agg_cell: float | str = round(agg, 2) if agg is not None else "[VALIDATE]"
      link_list = []
      for row in (cons_row, mod_row, agg_row):
        if row.get("source") == "WEB":
          for link in row.get("source_links", []):
            if link not in link_list:
              link_list.append(link)
      ws2.append([
        line,
        round(aws_val, 2),
        cons_cell,
        cons_row.get("source", "API"),
        mod_cell,
        mod_row.get("source", "API"),
        agg_cell,
        agg_row.get("source", "API"),
        "\n".join(link_list),
      ])
      current_row = ws2.max_row
      source_col_map = {4: cons_row.get("source", "API"), 6: mod_row.get("source", "API"), 8: agg_row.get("source", "API")}
      for col, source in source_col_map.items():
        fill = {
          "API": PatternFill(fill_type="solid", fgColor="C6EFCE"),
          "WEB": PatternFill(fill_type="solid", fgColor="FFF2CC"),
          "CACHE": PatternFill(fill_type="solid", fgColor="FCE4D6"),
          "NONE": PatternFill(fill_type="solid", fgColor="F8CBAD"),
        }.get(str(source), PatternFill(fill_type="solid", fgColor="D9E1F2"))
        ws2.cell(row=current_row, column=col).fill = fill
      totals["aws"] += aws_val
      if cons is None:
        scenario_has_validate["conservative"] = True
        scenario_unpriced_counts["conservative"] += 1
      else:
        totals["conservative"] += cons
      if mod is None:
        scenario_has_validate["moderate"] = True
        scenario_unpriced_counts["moderate"] += 1
      else:
        totals["moderate"] += mod
      if agg is None:
        scenario_has_validate["aggressive"] = True
        scenario_unpriced_counts["aggressive"] += 1
      else:
        totals["aggressive"] += agg

    total_cons: float | str = "N/A ([VALIDATE])" if scenario_has_validate["conservative"] else round(totals["conservative"], 2)
    total_mod: float | str = "N/A ([VALIDATE])" if scenario_has_validate["moderate"] else round(totals["moderate"], 2)
    total_agg: float | str = "N/A ([VALIDATE])" if scenario_has_validate["aggressive"] else round(totals["aggressive"], 2)

    ws2.append(["TOTAL", round(totals["aws"], 2), total_cons, "", total_mod, "", total_agg, "", ""])

    sav_cons: float | str = "N/A ([VALIDATE])" if scenario_has_validate["conservative"] else round(totals["aws"] - totals["conservative"], 2)
    sav_mod: float | str = "N/A ([VALIDATE])" if scenario_has_validate["moderate"] else round(totals["aws"] - totals["moderate"], 2)
    sav_agg: float | str = "N/A ([VALIDATE])" if scenario_has_validate["aggressive"] else round(totals["aws"] - totals["aggressive"], 2)
    ws2.append(["Savings $", "", sav_cons, "", sav_mod, "", sav_agg, "", ""])
    ws2.append([
      "Savings %",
      "",
      "N/A ([VALIDATE])" if scenario_has_validate["conservative"] else round(((totals["aws"] - totals["conservative"]) / totals["aws"] * 100) if totals["aws"] else 0, 2),
      "",
      "N/A ([VALIDATE])" if scenario_has_validate["moderate"] else round(((totals["aws"] - totals["moderate"]) / totals["aws"] * 100) if totals["aws"] else 0, 2),
      "",
      "N/A ([VALIDATE])" if scenario_has_validate["aggressive"] else round(((totals["aws"] - totals["aggressive"]) / totals["aws"] * 100) if totals["aws"] else 0, 2),
      "",
      "",
    ])
  else:
    totals = {"aws": 0.0, "conservative": 0.0, "moderate": 0.0, "aggressive": 0.0}
    scenario_has_validate = {"conservative": False, "moderate": False, "aggressive": False}
    scenario_unpriced_counts = {"conservative": 0, "moderate": 0, "aggressive": 0}
    ws2.append(["Line", "AWS Monthly", "Azure Monthly", "Savings $", "Savings %"])
    total_aws = 0.0
    total_az = 0.0
    for r in comparison_rows:
      aws_val = r["aws_monthly"]
      az_val = r["azure_monthly"]
      _register_line_value(r["line"], scenario, az_val)
      ws2.append([
        r["line"],
        round(aws_val, 2),
        "[VALIDATE]" if az_val is None else round(az_val, 2),
        "[VALIDATE]" if az_val is None else round(aws_val - az_val, 2),
        "[VALIDATE]" if az_val is None else round(((aws_val - az_val) / aws_val * 100) if aws_val else 0, 2),
      ])
      total_aws += aws_val
      if az_val is None:
        scenario_has_validate[scenario] = True
        scenario_unpriced_counts[scenario] += 1
      else:
        total_az += az_val
    totals["aws"] = total_aws
    totals[scenario] = total_az
    ws2.append([
      "TOTAL",
      round(total_aws, 2),
      "N/A ([VALIDATE])" if scenario_has_validate[scenario] else round(total_az, 2),
      "N/A ([VALIDATE])" if scenario_has_validate[scenario] else round(total_aws - total_az, 2),
      "N/A ([VALIDATE])" if scenario_has_validate[scenario] else round(((total_aws - total_az) / total_aws * 100) if total_aws else 0, 2),
    ])

  ws3 = wb.create_sheet("Assumptions & Levers")
  ws3.append(["Category", "Detail"])
  ws3.append(["Pricing source", "Azure Retail Prices API"])
  ws3.append(["Pricing date", pricing_date])
  ws3.append(["Region", region])
  ws3.append(["Currency", currency])
  ws3.append(["Scenario", scenario])
  ws3.append(["Levers", "conservative: PAYG only; moderate: 1Y Reservation + AHB + tiering; aggressive: 3Y Reservation + AHB + non-prod Dev/Test + Spot for interruptible only"])
  ws3.append(["Parity gaps", "; ".join(normalized.get("missing_or_ambiguous", [])) or "None captured"])
  ws3.append(["Open questions", "Confirm RTO/RPO, growth assumptions, non-prod interruptibility before final quote"])

  # Executive KPI sheet for polished stakeholder readout.
  ws4 = wb.create_sheet("Executive Summary")
  ws4.append(["Metric", "Value"])
  ws4.append(["Region", region])
  ws4.append(["Currency", currency])
  ws4.append(["Pricing date", pricing_date])
  ws4.append(["Input rows", len(aws_rows)])
  ws4.append(["Rows requiring validation", validate_count])

  available_scenarios = [s for s in SCENARIOS if not scenario_has_validate.get(s, False) and totals.get(s, 0.0) > 0]
  recommended_scenario = ""
  if available_scenarios:
    recommended_scenario = min(available_scenarios, key=lambda s: totals[s])

  for scenario_name in SCENARIOS:
    if scenario == "compare-all" or scenario == scenario_name:
      if scenario_has_validate.get(scenario_name, False):
        ws4.append([f"{scenario_name} azure monthly", "N/A ([VALIDATE])"])
        ws4.append([f"{scenario_name} savings vs AWS", "N/A ([VALIDATE])"])
        ws4.append([
          f"{scenario_name} priced subtotal monthly",
          f"{round(totals.get(scenario_name, 0.0), 2)} (excludes {scenario_unpriced_counts.get(scenario_name, 0)} line(s) pending validation)",
        ])
      else:
        savings = totals["aws"] - totals[scenario_name]
        savings_pct = (savings / totals["aws"] * 100.0) if totals["aws"] else 0.0
        ws4.append([f"{scenario_name} azure monthly", round(totals[scenario_name], 2)])
        ws4.append([f"{scenario_name} savings vs AWS", f"{round(savings, 2)} ({round(savings_pct, 2)}%)"])

  ws4.append(["Recommended scenario", recommended_scenario or "N/A (validation blockers)"])
  ws4.append(["Decision confidence", "High" if validate_count == 0 else ("Medium" if validate_count <= 5 else "Low")])

  # Validation sheet for auditability.
  ws5 = wb.create_sheet("Validation & Coverage")
  ws5.append(["Type", "Line", "Scenario", "Reason", "Service", "Source SKU", "Mapped SKU"])
  for item in validation_rows:
    ws5.append([
      "PRICE_LOOKUP",
      item["line"],
      item["scenario"],
      item["reason"],
      item["service"],
      item["source_sku"],
      item["mapped_sku"],
    ])
  missing_fields = normalized.get("missing_or_ambiguous", [])
  if missing_fields:
    for field in missing_fields:
      ws5.append(["INPUT_GAP", "", "", f"Missing or ambiguous spec field: {field}", "", "", ""])
  if not validation_rows and not missing_fields:
    ws5.append(["INFO", "", "", "No validation blockers detected.", "", "", ""])

  driver_scenario = recommended_scenario if recommended_scenario else (scenario if scenario in SCENARIOS else "conservative")
  driver_rows: List[Dict[str, Any]] = []
  for line, values in line_scenario_values.items():
    az = values.get(driver_scenario)
    aws_val = next((r["aws_monthly"] for r in comparison_rows if r["line"] == line), 0.0)
    if az is None:
      continue
    driver_rows.append({
      "line": line,
      "aws_monthly": aws_val,
      "azure_monthly": az,
      "delta": aws_val - az,
    })
  top_positive = sorted(driver_rows, key=lambda x: x["delta"], reverse=True)[:5]
  top_negative = sorted(driver_rows, key=lambda x: x["delta"])[:5]
  pricing_telemetry = _build_pricing_metrics(comparison_rows)

  output_path.parent.mkdir(parents=True, exist_ok=True)
  wb.save(output_path)

  scenario_totals: Dict[str, Any] = {"aws": round(totals["aws"], 2)}
  scenario_savings: Dict[str, Any] = {}
  priced_subtotals: Dict[str, Any] = {}
  for scenario_name in SCENARIOS:
    if scenario == "compare-all" or scenario == scenario_name:
      priced_subtotals[scenario_name] = {
        "amount": round(totals.get(scenario_name, 0.0), 2),
        "unpricedLines": scenario_unpriced_counts.get(scenario_name, 0),
      }
      if scenario_has_validate.get(scenario_name, False):
        scenario_totals[scenario_name] = "N/A ([VALIDATE])"
        scenario_savings[scenario_name] = "N/A ([VALIDATE])"
      else:
        azure_total = round(totals.get(scenario_name, 0.0), 2)
        savings = round(totals["aws"] - totals.get(scenario_name, 0.0), 2)
        savings_pct = round(((totals["aws"] - totals.get(scenario_name, 0.0)) / totals["aws"] * 100.0) if totals["aws"] else 0.0, 2)
        scenario_totals[scenario_name] = azure_total
        scenario_savings[scenario_name] = {"amount": savings, "percent": savings_pct}

  summary = {
    "region": region,
    "currency": currency,
    "pricingDate": pricing_date,
    "scenario": scenario,
    "rows": len(aws_rows),
    "validateCount": validate_count,
    "scenarioTotalsMonthly": scenario_totals,
    "scenarioPricedSubtotalsMonthly": priced_subtotals,
    "scenarioSavings": scenario_savings,
    "recommendedScenario": recommended_scenario or "N/A",
    "decisionConfidence": "High" if validate_count == 0 else ("Medium" if validate_count <= 5 else "Low"),
    "topDriversPositive": [
      {
        "line": d["line"],
        "aws": round(d["aws_monthly"], 2),
        "azure": round(d["azure_monthly"], 2),
        "delta": round(d["delta"], 2),
      }
      for d in top_positive
    ],
    "topDriversNegative": [
      {
        "line": d["line"],
        "aws": round(d["aws_monthly"], 2),
        "azure": round(d["azure_monthly"], 2),
        "delta": round(d["delta"], 2),
      }
      for d in top_negative
    ],
    "validationItemCount": len(validation_rows),
    "validationItems": validation_rows[:20],
    "validationItemsTruncated": len(validation_rows) > 20,
    "missingSpecFields": missing_fields,
    **pricing_telemetry,
  }
  return summary


def main() -> None:
  parser = argparse.ArgumentParser(description="Build Azure BOQ workbook from normalized input")
  parser.add_argument("--normalized", required=True)
  parser.add_argument("--output", required=True)
  parser.add_argument("--scenario", default="compare-all", choices=["compare-all", "conservative", "moderate", "aggressive"])
  parser.add_argument("--region", default="eastus")
  parser.add_argument("--currency", default="USD")
  parser.add_argument("--pricing-date", default="")
  args = parser.parse_args()

  normalized = json.loads(Path(args.normalized).read_text(encoding="utf-8"))
  pricing_meta = {
    "region": args.region,
    "currency": args.currency,
    "pricingDate": args.pricing_date,
  }
  summary = build_workbook(normalized, pricing_meta, Path(args.output), args.scenario)
  print(json.dumps(summary, indent=2))


if __name__ == "__main__":
  main()
